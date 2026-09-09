"""Dedup cascade, PRISMA flow and PRISMA-S appendix: edge cases and invariants."""
import csv

import pytest

from corpusslr import (Corpus, PrismaFlow, Record, deduplicate,
                       is_identifier_only_title, prisma_s_appendix,
                       prisma_s_markdown)


def _c(records_by_db):
    c = Corpus()
    for db, recs in records_by_db.items():
        c.add_records(recs, database=db, interface="API",
                      query=f"{db} query", date_run="2026-08-01")
    return c


# --------------------------------------------------------- invariants
def test_dedup_counts_are_internally_consistent():
    """sum(by_method) must equal before-after, else PRISMA-S misreports.

    Two works, each retrieved twice: the DOIs partition them, so the fuzzy
    stage must not collapse the two groups into one.
    """
    recs = [
        Record(title="Same paper on AI", doi="10.5001/a", year=2024,
               authors=["Smith, A"], source="Scopus"),
        Record(title="Same paper on AI", doi="10.5001/a", year=2024,
               authors=["Smith, A"], source="OpenAlex"),
        Record(title="Same paper on AI!", doi="10.5001/b", year=2024,
               authors=["Smith, A"], source="WoS"),
        Record(title="Same paper on AI.", doi="10.5001/b", year=2024,
               authors=["Smith, A"], source="PubMed"),
    ]
    res = deduplicate(recs)
    assert res.report.after == 2
    assert res.report.by_method == {"doi": 2}
    assert sum(res.report.by_method.values()) == res.report.removed


def test_uid_and_provenance_assigned_by_corpus():
    c = _c({"Scopus": [Record(title="T1", doi="10.5001/1"),
                       Record(title="T2", doi="10.5001/2")],
            "PubMed": [Record(title="T3", pmid="9")]})
    assert [r.uid for r in c.records] == ["R000001", "R000002", "R000003"]
    assert [r.search_id for r in c.records] == ["S1", "S1", "S2"]
    assert c.records[0].source == "Scopus"
    assert c.records[0].provenance[0]["database"] == "Scopus"
    assert c.get("R000002").title == "T2" and c.get("nope") is None
    assert len(c) == 3 and list(iter(c))[0] is c.records[0]


def test_no_duplicates_leaves_corpus_untouched():
    titles = ["Machine learning for credit risk",
              "Consumer trust in chatbots",
              "Supply chain resilience after shocks",
              "Genomic screening in primary care",
              "Urban heat islands and mortality"]
    recs = [Record(title=t, doi=f"10.5001/{i}", year=2020 + i, source="S")
            for i, t in enumerate(titles)]
    res = deduplicate(recs)
    assert res.report.after == 5 and res.report.removed == 0
    assert res.report.by_method == {} and res.report.overlap == {}
    assert res.report.overlap_markdown().startswith("_No cross-source")


def test_empty_corpus_is_handled():
    res = deduplicate([])
    assert res.records == [] and res.report.before == res.report.after == 0
    assert "0 -> 0" in res.report.summary()


def test_records_without_any_identifier_fall_back_to_fuzzy():
    a = Record(title="Machine learning for credit scoring", year=2021,
               authors=["Kowalski, Jan"], source="A")
    b = Record(title="Machine learning for credit scoring.", year=2021,
               authors=["Jan Kowalski"], source="B")
    res = deduplicate([a, b])
    assert res.report.after == 1 and res.report.by_method["fuzzy"] == 1
    d = res.report.decisions[0]
    assert d.method == "fuzzy" and 0.9 <= d.score <= 1.0


def test_fuzzy_rejects_different_first_authors():
    a = Record(title="Exercise and depression in adults", year=2020,
               authors=["Smith, Anna"], source="A")
    b = Record(title="Exercise and depression in adults", year=2020,
               authors=["Kowalski, Jan"], source="B")
    assert deduplicate([a, b]).report.after == 2


def test_fuzzy_threshold_and_year_tolerance_are_configurable():
    a = Record(title="Deep learning for medical imaging", year=2020,
               authors=["X, Y"], source="A")
    b = Record(title="Deep learning for medical imagery", year=2023,
               authors=["X, Y"], source="B")
    assert deduplicate([a, b]).report.after == 2               # defaults reject
    loose = deduplicate([a, b], fuzzy_threshold=0.85, year_tolerance=5)
    assert loose.report.after == 1


def test_missing_year_or_author_is_treated_permissively():
    a = Record(title="Governance of algorithmic systems", year=None,
               authors=[], source="A")
    b = Record(title="Governance of algorithmic systems", year=2019,
               authors=["Doe, J"], source="B")
    assert deduplicate([a, b]).report.after == 1


def test_merge_keeps_richest_and_unions_identifiers():
    poor = Record(title="AI in SMEs", doi="10.5001/x", source="Scopus",
                  scopus_id="850", abstract="trunc")
    rich = Record(title="AI in SMEs", doi="10.5001/x", source="OpenAlex",
                  openalex_id="W7", pmid="123", journal="JBR", year=2024,
                  authors=["A, B", "C, D"], issn="1-2", volume="1",
                  pages="1-9", doc_type="article", language="en",
                  url="http://x", keywords=["k"], cited_by=5,
                  abstract="a considerably longer abstract " * 5)
    res = deduplicate([poor, rich])
    m = res.records[0]
    assert m.scopus_id == "850" and m.openalex_id == "W7" and m.pmid == "123"
    assert m.journal == "JBR" and m.cited_by == 5 and len(m.abstract) > 50
    assert res.report.overlap[("OpenAlex", "Scopus")] == 1


def test_dedup_report_csv_roundtrip(tmp_path):
    recs = [Record(title="T", doi="10.5001/x", source="A"),
            Record(title="T", doi="10.5001/x", source="B")]
    res = deduplicate(recs)
    p = res.report.to_csv(str(tmp_path / "d.csv"))
    rows = list(csv.DictReader(open(p, encoding="utf-8")))
    assert len(rows) == 1
    assert rows[0]["method"] == "doi" and rows[0]["score"] == "1.0000"
    assert rows[0]["key"] == "10.5001/x"


def test_overlap_matrix_is_symmetric_markdown():
    recs = [Record(title="T1", doi="10.5001/1", source="Scopus"),
            Record(title="T1", doi="10.5001/1", source="WoS"),
            Record(title="T2", doi="10.5001/2", source="Scopus"),
            Record(title="T2", doi="10.5001/2", source="PubMed")]
    md = deduplicate(recs).report.overlap_markdown()
    assert "|Scopus|" in md and "|WoS|" in md and "|PubMed|" in md
    assert md.count("-") >= 3   # diagonal blanked


def test_same_source_duplicates_are_not_counted_as_overlap():
    recs = [Record(title="T", doi="10.5001/x", source="Scopus"),
            Record(title="T", doi="10.5001/x", source="Scopus")]
    res = deduplicate(recs)
    assert res.report.after == 1 and res.report.overlap == {}


def test_transitive_identifier_matching_merges_all_three():
    """r3 shares a DOI with r1 and a PMID with r2 -> all three are one work."""
    r1 = Record(title="Alpha study of treatment", doi="10.5001/shared",
                source="Scopus")
    r2 = Record(title="Completely different wording here", pmid="777",
                source="PubMed")
    r3 = Record(title="Alpha study of treatment", doi="10.5001/shared",
                pmid="777", source="OpenAlex")
    assert deduplicate([r1, r2, r3]).report.after == 1


def test_fuzzy_must_not_merge_records_with_conflicting_dois():
    a = Record(title="Climate risk disclosure in banking sector part 1",
               doi="10.5001/aaa", year=2022, authors=["Smith, A"], source="Scopus")
    b = Record(title="Climate risk disclosure in banking sector part 2",
               doi="10.5001/bbb", year=2022, authors=["Smith, A"], source="WoS")
    assert deduplicate([a, b]).report.after == 2


def test_fuzzy_must_not_merge_records_with_conflicting_pmids():
    a = Record(title="Randomised trial of drug X in adults", pmid="111",
               year=2021, authors=["Doe, J"], source="P1")
    b = Record(title="Randomised trial of drug Y in adults", pmid="222",
               year=2021, authors=["Doe, J"], source="P2")
    assert deduplicate([a, b]).report.after == 2


def test_decision_log_is_traceable_for_bare_record_lists():
    recs = [Record(title="T", doi="10.5001/x", source="A"),
            Record(title="T", doi="10.5001/x", source="B")]
    d = deduplicate(recs).report.decisions[0]
    assert d.kept_uid and d.removed_uid


def test_fuzzy_catches_duplicate_differing_in_first_characters():
    a = Record(title="The effect of artificial intelligence on productivity",
               year=2022, authors=["Smith, A"], source="A")
    b = Record(title="Effect of artificial intelligence on productivity",
               year=2022, authors=["Smith, A"], source="B")
    assert deduplicate([a, b]).report.after == 1


# --------------------------------------------------------- PRISMA flow
def test_flow_counts_dict_and_markdown():
    c = _c({"Scopus": [Record(title=f"T{i}", doi=f"10.5001/{i}") for i in range(10)],
            "PubMed": [Record(title="T0", doi="10.5001/0", pmid="1")]})
    res = deduplicate(c)
    flow = PrismaFlow.from_dedup(c, res).set_screening(
        records_excluded=4, reports_not_retrieved=1,
        fulltext_exclusions={"wrong population": 2, "not empirical": 1},
        studies_included=2, reports_included=3,
        automation_excluded=0, other_excluded=0)
    cnt = flow.counts()
    assert cnt["identified"] == 11 and cnt["duplicates_removed"] == 1
    assert cnt["records_screened"] == 10 and cnt["reports_sought"] == 6
    assert cnt["reports_assessed"] == 5 and cnt["studies_included"] == 2
    assert flow.validate() == []
    md = flow.to_markdown()
    assert "Records identified from databases (n = 11)" in md
    assert "wrong population (n = 2)" in md
    assert "Reports of included studies (n = 3)" in md


def test_flow_detects_each_arithmetic_violation():
    assert PrismaFlow(db_counts={"A": 3}, duplicates_removed=5).validate()
    assert PrismaFlow(db_counts={"A": 3}, records_excluded=5).validate()
    assert PrismaFlow(db_counts={"A": 3}, reports_not_retrieved=5).validate()
    bad = PrismaFlow(db_counts={"A": 10}, records_excluded=5,
                     fulltext_exclusions={"r": 1}, studies_included=99)
    assert any("studies included" in p for p in bad.validate())


def test_svg_escapes_markup_and_renders_all_boxes(tmp_path):
    flow = PrismaFlow(db_counts={"Scopus & <WoS>": 5}, duplicates_removed=1,
                      dedup_by_method={"doi": 1})
    flow.set_screening(records_excluded=1, fulltext_exclusions={"a<b": 1},
                       studies_included=3)
    svg = flow.to_svg(str(tmp_path / "f.svg"))
    assert svg.startswith("<svg") and svg.rstrip().endswith("</svg>")
    assert "&amp;" in svg and "&lt;WoS&gt;" in svg
    assert "Scopus & <WoS>" not in svg          # raw markup never leaks
    for label in ("Identification", "Screening", "Included"):
        assert f">{label}<" in svg
    assert "[doi: 1]" in svg
    assert svg.count("<rect") >= 9              # 8 boxes + background


def test_svg_without_fulltext_exclusions_still_renders_zero_box():
    svg = PrismaFlow(db_counts={"A": 2}).to_svg()
    assert "Reports excluded (n = 0)" in svg


def test_validate_catches_impossible_flow_with_zero_included():
    flow = PrismaFlow(db_counts={"A": 10}, records_excluded=8)
    flow.set_screening(records_excluded=8,
                       fulltext_exclusions={"reason": 50},
                       studies_included=0)
    assert flow.validate()


def test_negative_screening_numbers_are_rejected():
    flow = PrismaFlow(db_counts={"A": 10})
    flow.set_screening(records_excluded=-5, studies_included=0)
    assert flow.validate()


def test_validate_requires_fulltext_balance_even_at_zero_included():
    """A review including no studies must still account for what it assessed."""
    flow = PrismaFlow(db_counts={"A": 10}).set_screening(
        records_excluded=8, fulltext_exclusions={"reason": 1},
        studies_included=0)
    assert any("studies included" in p for p in flow.validate())
    ok = PrismaFlow(db_counts={"A": 10}).set_screening(
        records_excluded=8, fulltext_exclusions={"reason": 2},
        studies_included=0)
    assert ok.validate() == []


@pytest.mark.parametrize("kwargs", [
    {"records_excluded": -5, "studies_included": 0},
    {"records_excluded": 0, "reports_not_retrieved": -1, "studies_included": 0},
    {"records_excluded": 0, "studies_included": -2},
    {"records_excluded": 0, "fulltext_exclusions": {"r": -3},
     "studies_included": 0},
])
def test_negative_counts_are_reported(kwargs):
    flow = PrismaFlow(db_counts={"A": 10}).set_screening(**kwargs)
    assert any("Negative" in p for p in flow.validate())


def test_negative_source_count_is_reported():
    assert any("Negative record count" in p
               for p in PrismaFlow(db_counts={"A": -1}).validate())


def test_dedup_assigns_traceable_uids_to_bare_lists():
    recs = [Record(title="T", doi="10.5001/x", source="A"),
            Record(title="T", doi="10.5001/x", source="B")]
    d = deduplicate(recs).report.decisions[0]
    assert d.kept_uid and d.removed_uid and d.kept_uid != d.removed_uid


def test_dedup_preserves_existing_corpus_uids():
    c = _c({"Scopus": [Record(title="T", doi="10.5001/x")],
            "PubMed": [Record(title="T", doi="10.5001/x")]})
    d = deduplicate(c).report.decisions[0]
    assert d.kept_uid == "R000001" and d.removed_uid == "R000002"


# ------------------------------------------------------------ PRISMA-S
def test_prisma_s_reports_every_search_event_and_totals():
    c = _c({"Scopus": [Record(title="T1", doi="10.5001/1")],
            "Embase": [Record(title="T2", doi="10.5001/2")]})
    c.searches[1].filters = "years 2015-2026"
    c.searches[1].url = "https://api.example.org"
    c.searches[1].notes = "vendor export"
    res = deduplicate(c)
    md = prisma_s_markdown(c, res)
    assert "| S1 | Scopus |" in md and "| S2 | Embase |" in md
    assert "Total records identified: **2**" in md
    assert "years 2015-2026" in md and "api.example.org" in md
    assert "vendor export" in md
    assert "Item 16" in md


def test_prisma_s_without_dedup_result_and_unrecorded_query():
    c = Corpus()
    c.add_records([Record(title="T")], database="Scopus")
    md = prisma_s_markdown(c, None)
    assert "Deduplication was not performed" in md
    assert "(not recorded)" in md


def test_prisma_s_appendix_writes_markdown(tmp_path):
    c = _c({"Scopus": [Record(title="T", doi="10.5001/x")]})
    p = prisma_s_appendix(c, deduplicate(c), str(tmp_path / "app.md"))
    assert open(p, encoding="utf-8").read().startswith("# Search strategy")


def test_prisma_s_appendix_docx(tmp_path):
    pytest.importorskip("docx")
    c = _c({"Scopus": [Record(title="T", doi="10.5001/x", year=2024)],
            "PubMed": [Record(title="T", doi="10.5001/x", pmid="5")]})
    p = prisma_s_appendix(c, deduplicate(c), str(tmp_path / "app.docx"))
    assert p.endswith(".docx")
    from docx import Document
    doc = Document(p)
    text = "\n".join(par.text for par in doc.paragraphs)
    assert "PRISMA-S" in text
    assert len(doc.tables) >= 1
    assert doc.tables[0].cell(0, 0).text == "#"


def test_prisma_s_reports_custom_thresholds():
    c = _c({"A": [Record(title="T", doi="10.5001/x")]})
    md = prisma_s_markdown(c, deduplicate(c), fuzzy_threshold=0.8,
                           year_tolerance=3)
    assert "similarity >= 0.8" in md and "+/-3" in md


# ------------------------------------------------- survivor determinism
def test_surviving_record_is_the_first_encountered():
    """The kept uid must not depend on which export was most complete.

    merge_from() fills every blank from the removed copies, so completeness is
    identical either way; keeping the first-encountered copy is what a reference
    manager does and is what makes record_id stable across runs.
    """
    a = Record(title="Shared paper", doi="10.5001/x", source="Scopus", uid="R000001")
    b = Record(title="Shared paper", doi="10.5001/x", source="Embase", uid="R000002",
               abstract="A far longer abstract that makes this copy richer.",
               pages="10-20", volume="7", issue="3")
    res = deduplicate([a, b])
    assert [r.uid for r in res.records] == ["R000001"]
    kept = res.records[0]
    assert kept.abstract == b.abstract     # richer metadata still merged in
    assert (kept.pages, kept.volume, kept.issue) == ("10-20", "7", "3")


def test_survivor_is_stable_under_input_permutation_of_removed_copies():
    """Reordering the duplicate copies must not change which uid survives."""
    def run(order):
        recs = [Record(title="Same work here", doi="10.5001/y", uid=f"R{i:06d}",
                       source=s, abstract="x" * n)
                for i, (s, n) in enumerate(order, 1)]
        return [r.uid for r in deduplicate(recs).records]
    assert run([("A", 10), ("B", 400), ("C", 100)]) == ["R000001"]
    assert run([("A", 400), ("B", 10), ("C", 100)]) == ["R000001"]


def test_conflicting_journal_titles_resolve_by_majority():
    """A single mislabelled export must not impose its journal on the cluster."""
    recs = [
        Record(title="AI adoption in small firms", doi="10.5001/z",
               journal="IEEE Access", source="IEEE"),
        Record(title="AI adoption in small firms", doi="10.5001/z",
               journal="Journal of Business Research", source="Scopus"),
        Record(title="AI adoption in small firms", doi="10.5001/z",
               journal="Journal of Business Research", source="Embase"),
    ]
    assert deduplicate(recs).records[0].journal == "Journal of Business Research"


def test_pairwise_disagreement_keeps_the_survivors_value():
    """With only two copies there is no majority, so nothing is overridden."""
    recs = [Record(title="T", doi="10.5001/p", journal="Journal A", source="X"),
            Record(title="T", doi="10.5001/p", journal="Journal B", source="Y")]
    assert deduplicate(recs).records[0].journal == "Journal A"


# ------------------------------------------- locus guard & blocking rounds
def test_supplement_doi_split_by_first_page():
    """One DOI over a whole supplement must not merge abstracts on different pages."""
    recs = [
        Record(title="Effects of drug A on vascular outcomes", doi="10.5001/suppl",
               year=2020, pages="S12", volume="63", authors=["Smith, A"],
               source="Embase"),
        Record(title="Effects of drug A on vascular outcomes", doi="10.5001/suppl",
               year=2020, pages="S12", volume="63", authors=["Smith, A"],
               source="Scopus"),
        Record(title="Effects of drug B on renal outcomes", doi="10.5001/suppl",
               year=2020, pages="S48", volume="63", authors=["Brown, C"],
               source="Embase"),
    ]
    res = deduplicate(recs)
    assert res.report.after == 2
    assert res.report.by_locus.get("pages") == 1
    assert any("titles disagreed" in w or "rejected" in w
               for w in res.report.warnings())


def test_locus_guard_can_be_disabled():
    recs = [Record(title="Same abstract text", doi="10.5001/s", year=2020,
                   pages=p, volume="63", authors=["Smith, A"], source=s)
            for p, s in (("S12", "A"), ("S12", "B"), ("S48", "C"))]
    assert deduplicate(recs, separate_by_locus=False).report.after == 1


def test_locus_guard_needs_both_values_present():
    """A missing page must never be read as disagreement."""
    recs = [Record(title="Shared work", doi="10.5001/x", year=2020, pages="S12",
                   volume="63", authors=["Smith, A"], source="A"),
            Record(title="Shared work", doi="10.5001/x", year=2020, pages="",
                   volume="63", authors=["Smith, A"], source="B"),
            Record(title="Shared work", doi="10.5001/x", year=2020, pages="S12",
                   volume="", authors=["Smith, A"], source="C")]
    assert deduplicate(recs).report.after == 1


def test_volume_guard_requires_the_year_to_disagree_too():
    """Databases render volumes inconsistently, so volume alone cannot separate."""
    recs = [Record(title="One work", doi="10.5001/v", year=2020, volume="63",
                   authors=["Smith, A"], source="A"),
            Record(title="One work", doi="10.5001/v", year=2020, volume="7",
                   authors=["Smith, A"], source="B"),
            Record(title="One work", doi="10.5001/v", year=2020, volume="63",
                   authors=["Smith, A"], source="C")]
    assert deduplicate(recs).report.after == 1


def test_blocking_round_recovers_duplicate_with_truncated_title():
    """Export truncation drops title similarity below the fuzzy threshold, but
    author + year + first page still identify the record."""
    recs = [
        Record(title="Long-term outcomes of statin therapy in elderly patients",
               authors=["Nowak, J"], year=2019, pages="1123-1130", volume="41",
               journal="European Heart Journal", source="WoS"),
        Record(title="Long-term outcomes of statin ther...",
               authors=["Nowak, J"], year=2019, pages="1123", volume="41",
               journal="European Heart Journal", source="ProQuest"),
    ]
    assert deduplicate(recs).report.after == 1
    assert deduplicate(recs, blocking_rounds=False).report.after == 2


def test_blocking_rounds_do_not_merge_records_at_different_loci():
    """A composite key must not override the locus guard."""
    recs = [Record(title="Trial of therapy X", authors=["Lee, K"], year=2021,
                   pages="10", volume="5", journal="J Med", doi="10.5001/a",
                   source="A"),
            Record(title="Trial of therapy X", authors=["Lee, K"], year=2021,
                   pages="88", volume="5", journal="J Med", doi="10.5001/b",
                   source="B")]
    assert deduplicate(recs).report.after == 2


def test_page_and_volume_normalization():
    from corpusslr.dedup import normalize_pages, normalize_volume
    assert normalize_pages("413-420") == normalize_pages("413--420") == "413"
    assert normalize_pages("S354") == "s354"
    assert normalize_pages("pp. 77") == "77"
    assert normalize_pages("") == "" and normalize_pages(None) == ""
    assert normalize_volume("38 (Supplement 1)") == "38"
    assert normalize_volume("3)") == "3"
    assert normalize_volume("") == ""


def test_report_exposes_round_and_locus_diagnostics():
    recs = [Record(title="Effects of drug A here", doi="10.5001/s", year=2020,
                   pages=p, volume="63", authors=["Smith, A"], source=s)
            for p, s in (("S12", "A"), ("S12", "B"), ("S48", "C"))]
    rep = deduplicate(recs).report
    assert rep.by_locus and rep.id_links_rejected >= 1
    assert isinstance(rep.round_candidates, int)


# ---------------------------------------------- normalization edge cases
def test_percent_encoded_doi_matches_its_decoded_form():
    """A DOI serialised out of a URL must not become a second identity."""
    a = Record(title="Statin therapy outcomes", doi="10.1016/s0025-7753%2817%2930624-3",
               year=2017, source="Embase")
    b = Record(title="Statin therapy outcomes", doi="10.1016/s0025-7753(17)30624-3",
               year=2017, source="Scopus")
    assert a.doi == b.doi
    assert deduplicate([a, b]).report.after == 1


def test_double_encoded_doi_is_decoded():
    from corpusslr import normalize_doi
    assert normalize_doi("10.1016/x%252817%2529y") == "10.1016/x(17)y"
    assert normalize_doi("10.5001/plain") == "10.5001/plain"


def test_excel_mangled_pages_never_block_a_merge():
    """Excel turns '11-9' into '11-Sep' and '11-19' into 'Nov-19'; the two
    renderings of one range must not read as different loci."""
    from corpusslr.dedup import excel_mangled_pages
    assert excel_mangled_pages("Nov-19") and excel_mangled_pages("11-Sep")
    assert not excel_mangled_pages("413-420") and not excel_mangled_pages("S354")
    assert not excel_mangled_pages("") and not excel_mangled_pages(None)

    recs = [Record(title="Vascular effects of incretin", doi="10.5001/shared",
                   year=2015, pages=p, volume="73", authors=["Mita, T"], source=s)
            for p, s in (("Nov-19", "A"), ("11-Sep", "B"), ("Nov-19", "C"))]
    assert deduplicate(recs).report.after == 1


def test_article_number_padding_does_not_separate_records():
    """'137960' and 'e0137960' are the same PLoS ONE article number."""
    recs = [Record(title="Exendin-4 prevents proliferation", doi="10.1371/journal.pone.0137960",
                   year=2015, pages=p, volume="10", authors=["Nagayama, K"], source=s)
            for p, s in (("137960", "A"), ("e0137960", "B"), ("137960", "C"))]
    assert deduplicate(recs).report.after == 1


# ------------------------------------------------- co-publication override
def test_two_publisher_dois_for_one_work_are_merged():
    """One conference abstract printed in two journals of the same publisher
    carries two DOIs; total agreement elsewhere overrides the conflict."""
    a = Record(title="Liraglutide and linagliptin improve glycemia",
               doi="10.1016/j.regpep.2012.05.039", year=2012, pages="e24",
               volume="1", journal="Appetite", authors=["Smith, A"], source="A")
    b = Record(title="Liraglutide and linagliptin improve glycemia",
               doi="10.1016/j.appet.2012.05.003", year=2012, pages="e24",
               volume="59", journal="Appetite", authors=["Smith, A"], source="B")
    res = deduplicate([a, b])
    assert res.report.after == 1
    assert res.report.copublication_merges == 1
    assert deduplicate([a, b], allow_copublication=False).report.after == 2


def test_copublication_override_still_blocks_genuinely_distinct_works():
    """The override must not reopen the Part 1 / Part 2 false positive."""
    a = Record(title="Climate risk disclosure in banking part 1", doi="10.5001/aaa",
               year=2022, pages="10", authors=["Smith, A"], source="Scopus")
    b = Record(title="Climate risk disclosure in banking part 2", doi="10.5001/bbb",
               year=2022, pages="10", authors=["Smith, A"], source="WoS")
    assert deduplicate([a, b]).report.after == 2
    # differing first page also keeps the block closed
    c = Record(title="Trial of drug X in adults", doi="10.5001/ccc", year=2021,
               pages="10", authors=["Doe, J"], source="A")
    d = Record(title="Trial of drug X in adults", doi="10.5001/ddd", year=2021,
               pages="88", authors=["Doe, J"], source="B")
    assert deduplicate([c, d]).report.after == 2


# ------------------------------------------- conference vs journal article
def test_conference_abstract_kept_separate_from_the_journal_article():
    """Cochrane and ASySD treat these as two reports of one study."""
    from corpusslr.dedup import is_conference_venue
    art = Record(title="Increased DPP-4 accelerates vascular aging",
                 doi="10.1016/j.ijcard.2017.05.062", year=2017, pages="413-420",
                 volume="243", journal="Int J Cardiol", authors=["Lei, Y"],
                 source="Scopus")
    conf = Record(title="Increased DPP-4 accelerates vascular aging", doi="",
                  year=2017, pages="", volume="136",
                  journal="Circulation. Conference: Resuscitation Science",
                  authors=["Lei, Y"], source="Embase")
    assert is_conference_venue(conf) and not is_conference_venue(art)
    assert deduplicate([art, conf]).report.after == 2
    assert deduplicate([art, conf], separate_conference=False).report.after == 1


def test_conference_split_needs_missing_article_coordinates():
    """A conference record WITH a DOI and pages is an article-level record and
    must still merge with its twin."""
    a = Record(title="Effects of therapy X", doi="10.5001/x", year=2020, pages="S12",
               volume="63", journal="Diabetes. Conference: ADA", authors=["Lee, K"],
               source="A")
    b = Record(title="Effects of therapy X", doi="10.5001/x", year=2020, pages="S12",
               volume="63", journal="Diabetes", authors=["Lee, K"], source="B")
    assert deduplicate([a, b]).report.after == 1


def test_two_conference_records_of_one_abstract_still_merge():
    recs = [Record(title="Effects of therapy Y", doi="", year=2020, pages="",
                   volume="136", journal="Circulation. Conference: AHA",
                   authors=["Lee, K"], source=s) for s in ("Embase", "Scopus")]
    assert deduplicate(recs).report.after == 1


# --------------------------------------------------- length prefilter
def test_length_prefilter_never_rejects_a_reachable_pair():
    """The prefilter is a necessary condition, so it must not change any verdict.

    Brute-force check: for every pair in a mixed corpus, if the real ratio meets
    the threshold then the cheap bound must also permit it.
    """
    from corpusslr.dedup import _length_permits, _title_ratio
    titles = ["machine learning for credit risk assessment",
              "machine learning for credit risk", "credit risk",
              "deep learning applied to consumer credit scoring models",
              "a", "machine learning for credit risk assessment in banking",
              "consumer trust in conversational agents", ""]
    recs = [Record(title=t, source="S") for t in titles]
    for thr in (0.70, 0.85, 0.93, 0.99):
        for a in recs:
            for b in recs:
                if a is b:
                    continue
                if _title_ratio(a, b) >= thr:
                    assert _length_permits(a, b, thr), (a.title, b.title, thr)


def test_length_prefilter_rejects_hopeless_pairs():
    from corpusslr.dedup import _length_permits
    short = Record(title="ai", source="A")
    long = Record(title="artificial intelligence adoption in small firms", source="B")
    assert not _length_permits(short, long, 0.93)
    assert not _length_permits(Record(title="", source="A"), long, 0.5)


def test_prefilter_does_not_change_dedup_outcome():
    """A duplicate differing only by an added subtitle must still be found."""
    a = Record(title="Machine learning for credit risk assessment",
               year=2021, authors=["Nowak, J"], source="A")
    b = Record(title="Machine learning for credit risk assessment.",
               year=2021, authors=["Nowak, J"], source="B")
    assert deduplicate([a, b]).report.after == 1


# ------------------------------------------- pseudo page ranges (article length)
def test_pseudo_page_range_detection():
    from corpusslr.dedup import pseudo_page_range
    assert pseudo_page_range("1-15") and pseudo_page_range("1 - 8")
    assert not pseudo_page_range("1-1")        # a one-page article at page 1
    assert not pseudo_page_range("413-420")    # a real location
    assert not pseudo_page_range("e0137960")
    assert not pseudo_page_range("") and not pseudo_page_range(None)
    assert not pseudo_page_range("1-1500")     # too long to be an article length


def test_article_length_masquerading_as_pages_does_not_block_a_merge():
    """DOAJ stores an article's LENGTH in the page range for article-numbered
    journals, so one copy claims page 1 while the other carries article no. 393."""
    a = Record(title="Fostering green transformational leadership", year=2024,
               journal="BMC Nursing", volume="23", pages="1-15", source="DOAJ",
               authors=["Ali, M"])
    b = Record(title="Fostering green transformational leadership", year=2024,
               journal="BMC Nursing", volume="23", pages="393", source="Europe PMC",
               authors=["Ali, M"])
    assert deduplicate([a, b]).report.after == 1


def test_real_page_locations_still_separate_records():
    """The defence must not disarm the locus guard for genuine page ranges.

    The guard only fires on an identifier shared by three or more records -- the
    signature of a supplement-wide DOI -- so the fixture supplies three.
    """
    recs = [Record(title="Effects of drug A on outcomes", doi="10.5001/suppl",
                   year=2020, volume="63", pages=p, authors=["Smith, A"], source=s)
            for p, s in (("45-52", "A"), ("374-380", "B"), ("512-519", "C"))]
    assert deduplicate(recs).report.after == 3
    # and the same three merge once the pages are unreadable article lengths
    for r in recs:
        r.pages = "1-8"
    assert deduplicate(recs).report.after == 1


# --------------------------------------------------------------------------
# Missing-value literals in the DOI field.
#
# Exports write "no DOI" as a literal: R writes NA, several databases write
# N/A, NULL, none or a bare dash. If such a value survives normalization it
# becomes a shared identifier, and a shared identifier is the strongest
# evidence the cascade has -- so every record in the file merges with every
# other one, silently, while the PRISMA count still adds up. Measured on the
# ASySD Diabetes file, whose missing DOIs are the literal "NA": 492 distinct
# works shared the single key "na".
# --------------------------------------------------------------------------
@pytest.mark.parametrize("marker", ["NA", "N/A", "n/a", "NULL", "none", "-", "--"])
def test_missing_value_markers_never_link_two_records(marker):
    a = Record(title="Coronary microvascular function in diabetes",
               authors=["Smith, A"], year=2021, journal="Diabetologia", doi=marker)
    b = Record(title="Wind turbine blade fatigue under cyclic load",
               authors=["Kowalski, J"], year=2019, journal="Wind Energy", doi=marker)
    result = deduplicate([a, b])
    assert result.report.after == 2, (
        "{!r} in the DOI field must not act as an identifier".format(marker))


def test_a_whole_corpus_of_NA_dois_is_not_collapsed():
    """The failure mode is quadratic: it is the corpus, not the pair, that dies.

    The titles are lexically unrelated on purpose. An earlier version of this
    test numbered them ("Distinct study number 1", "... 2"), which differ by one
    character and are therefore genuine fuzzy matches -- the fixture, not the
    package, was wrong.
    """
    titles = ["Coronary microvascular function in type 2 diabetes",
              "Wind turbine blade fatigue under cyclic loading",
              "Lexical borrowing in medieval Castilian charters",
              "Photocatalytic degradation of azo dyes on titania",
              "Voter turnout and compulsory registration in Belgium",
              "Gut microbiome composition after antibiotic exposure",
              "Topological edge states in photonic crystals",
              "Last-mile delivery costs in dense urban networks",
              "Technical debt accumulation in microservice migrations",
              "Sea surface temperature anomalies in the Baltic",
              "Minimum wage effects on teenage employment",
              "CRISPR off-target detection by whole-genome sequencing"]
    recs = [Record(title=t, authors=["Author{}, X".format(i)], year=2000 + i,
                   journal="Journal of Field {}".format(i), doi="NA")
            for i, t in enumerate(titles)]
    assert deduplicate(recs).report.after == len(titles)


def test_a_real_doi_still_links_two_records():
    """The guard must not cost the cascade its actual job."""
    a = Record(title="Array programming with NumPy", year=2020,
               doi="10.1038/s41586-020-2649-2")
    b = Record(title="Array programming with NumPy", year=2020,
               doi="https://doi.org/10.1038/S41586-020-2649-2")
    assert deduplicate([a, b]).report.after == 1


# --------------------------------------------------------------------------
# Instalments of a recurring publication.
#
# Grey literature is full of series whose instalments share a title verbatim and
# differ only in the marker naming which instalment it is. Fuzzy matching cannot
# separate them because the titles genuinely are ~97% identical. Measured on a
# live transport corpus (electric-vehicle charging infrastructure), measured
# with the rule disabled and enabled: 19 false merges eliminated (39 -> 20 over
# 182 judgeable pairs), all 19 consecutive quarters of one U.S. Department of
# Energy report series. Merging them destroys a time series inside a review.
# --------------------------------------------------------------------------
_OSTI = ("Electric Vehicle Charging Infrastructure Trends from the "
         "Alternative Fueling Station Locator")


@pytest.mark.parametrize("title_a,title_b", [
    ("{} (Third Quarter 2021)".format(_OSTI), "{}: Fourth Quarter 2021".format(_OSTI)),
    ("Annual energy outlook 2019", "Annual energy outlook 2020"),
    ("National travel survey, wave 2", "National travel survey, wave 3"),
    ("Freight statistics, Vol. 3", "Freight statistics, Vol. 4"),
    ("Mobility report: January", "Mobility report: February"),
    ("Congestion study, Part 1", "Congestion study, Part 2"),
])
def test_consecutive_instalments_are_not_merged(title_a, title_b):
    a = Record(title=title_a, year=2021, journal="Agency series", doi="10.2172/1855378")
    b = Record(title=title_b, year=2021, journal="Agency series", doi="10.2172/1867218")
    assert deduplicate([a, b]).report.after == 2, (
        "{!r} and {!r} are different instalments".format(title_a, title_b))


@pytest.mark.parametrize("title", [
    "COVID-19 impact on urban transport",       # a year-like token in a disease name
    "Effects of the 2008 financial crisis",     # a year that is part of the subject
    "Q1 congestion patterns",                   # the same marker on both copies
    "Deep learning for diabetic retinopathy",   # no marker at all
])
def test_two_copies_of_one_work_still_merge(title):
    """The guard must fire on a DIFFERENCE in markers, never on their presence."""
    a = Record(title=title, year=2021, journal="J", volume="12", pages="1-9",
               doi="10.5001/same")
    b = Record(title=title, year=2021, journal="J", volume="12", pages="1-9", doi="")
    assert deduplicate([a, b]).report.after == 1


def test_the_separating_reason_is_reported_not_silent():
    """A reviewer must be able to see why two near-identical titles stayed apart."""
    from corpusslr.dedup import _separating_field
    a = Record(title="{} (Third Quarter 2021)".format(_OSTI), year=2021)
    b = Record(title="{}: Fourth Quarter 2021".format(_OSTI), year=2021)
    assert _separating_field(a, b) == "series instalment"


# --------------------------------------------------------------------------
# Titles that are only an identifier.
#
# Some publishers deposit the DOI, a handle or an ISBN into the title field.
# Two such "titles" from one journal issue differ only in their final digits
# and score ~0.95 on any string metric, so a matcher that trusts them merges
# distinct articles. Observed live: five Crossref records from one issue of
# *Journal of Social Development in Africa* whose title field is the literal
# DOI; with identifiers hidden, four separate papers collapsed into one.
# --------------------------------------------------------------------------
_DOI_TITLES = ["10.4314/jsda.v17i2.23834", "10.4314/jsda.v17i2.23837",
               "10.4314/jsda.v17i2.23840", "10.4314/jsda.v17i2.23842"]


def test_records_titled_only_by_their_doi_stay_apart():
    recs = [Record(title=t, year=2015, journal="J Soc Dev Africa")
            for t in _DOI_TITLES]
    assert deduplicate(recs).report.after == len(_DOI_TITLES)


@pytest.mark.parametrize("title", [
    "10.4314/jsda.v17i2.23834",
    "https://doi.org/10.1038/s41586-020-2649-2",
    "doi: 10.1016/j.softx.2025.102416",
    "hdl.handle.net/2027/uc1.b3305076",
    "ISBN 978-0-306-40615-7",
    "http://example.org/record/1234",
])
def test_identifier_only_titles_are_recognised(title):
    assert is_identifier_only_title(title)


@pytest.mark.parametrize("title", [
    "Deep learning for diabetic retinopathy",
    "10 years of climate policy in the European Union",   # leading number
    "Nature 10.5 million years ago",                      # decimal in prose
    "ISBN allocation policy in small presses",            # the word, not a number
    "A study of https://example.org as a citation target",  # URL inside a title
])
def test_real_titles_are_not_mistaken_for_identifiers(title):
    assert not is_identifier_only_title(title)


def test_a_real_duplicate_pair_still_merges_on_title():
    """The guard must cost nothing when titles carry actual content."""
    a = Record(title="Social capital and civic participation in rural Poland",
               year=2019, journal="Sociologia Ruralis")
    b = Record(title="Social capital and civic participation in rural Poland",
               year=2019, journal="Sociologia Ruralis")
    assert deduplicate([a, b]).report.after == 1


def test_the_identifier_is_not_thrown_away_with_the_title():
    """A DOI-titled record still deduplicates through the identifier stage."""
    a = Record(title="10.4314/jsda.v17i2.23834", doi="10.4314/jsda.v17i2.23834",
               year=2015)
    b = Record(title="Land reform and rural livelihoods", year=2015,
               doi="10.4314/jsda.v17i2.23834")
    assert deduplicate([a, b]).report.after == 1
